import openpyxl
import pymupdf
from docx import Document

from functions import (
    create_pdf,
    open_excel,
    prepare_data,
    read_docx,
    read_pdf,
    read_txt,
    remove_phone_numbers_docx,
    remove_phone_numbers_pdf,
    remove_phone_numbers_txt,
    )


def test_read_docx(tmp_path):
    input_file = tmp_path / "example.docx"

    doc = Document()
    doc.add_paragraph("John Smith, +1 234 567 8910")
    doc.add_paragraph("Alice Brown, +44 123 456 7890")
    doc.save(input_file)

    result = read_docx(input_file)

    assert "John Smith" in result
    assert "Alice Brown" in result


def test_read_docx_invalid_file():
    result = read_docx("not_found.docx")

    assert result.startswith("error:")


def test_read_txt(tmp_path):
    input_file = tmp_path / "example.txt"
    input_file.write_text("John Smith, +1 234 567 8910", encoding="utf-8")

    result = read_txt(input_file)

    assert result == "John Smith, +1 234 567 8910"


def test_read_txt_invalid_file():
    result = read_txt("not_found.txt")

    assert result.startswith("error:")


def test_read_pdf(tmp_path):
    input_file = tmp_path / "example.pdf"

    create_pdf("John Smith, +1 234 567 8910", input_file)

    result = read_pdf(input_file)

    assert "John Smith" in result
    assert "+1 234 567 8910" in result


def test_read_pdf_invalid_file():
    result = read_pdf("not_found.pdf")

    assert result.startswith("error:")


def test_prepare_data():
    text = """
    John Smith +1 234 567 8910
    Alice Brown +44 123 456 7890
    """

    result = prepare_data(text)

    assert result == [
        {"name": "John Smith", "Phone Number": "+1 234 567 8910"},
        {"name": "Alice Brown", "Phone Number": "+44 123 456 7890"}]


def test_prepare_data_without_country_code():
    text = "John Smith 234 567 8910"

    result = prepare_data(text)

    assert result == [
        {"name": "John Smith", "Phone Number": "234 567 8910"}]


def test_prepare_data_without_match():
    result = prepare_data("There is no phone number here.")

    assert result == []


def test_open_excel_create_new_file(tmp_path):
    excel_file = tmp_path / "result.xlsx"
    data = [
        {"name": "John Smith", "Phone Number": "+1 234 567 8910"},
        {"name": "Alice Brown", "Phone Number": "+44 123 456 7890"}]

    open_excel(excel_file, data)

    workbook = openpyxl.load_workbook(excel_file)
    sheet = workbook.active

    assert sheet["A1"].value == "Name"
    assert sheet["B1"].value == "Phone Number"
    assert sheet["A2"].value == "John Smith"
    assert sheet["B3"].value == "+44 123 456 7890"


def test_open_excel_update_existing_file(tmp_path):
    excel_file = tmp_path / "result.xlsx"
    data = [{"name": "John Smith", "Phone Number": "+1 234 567 8910"}]

    open_excel(excel_file, data)
    open_excel(excel_file, data)

    workbook = openpyxl.load_workbook(excel_file)
    sheet = workbook.active

    assert sheet.max_row == 3


def test_remove_phone_numbers_docx(tmp_path):
    input_file = tmp_path / "example.docx"
    output_file = tmp_path / "result.docx"

    doc = Document()
    doc.add_paragraph("John Smith, +1 234 567 8910")
    doc.add_paragraph("Alice Brown, +44 123 456 7890")
    doc.save(input_file)

    remove_phone_numbers_docx(input_file, output_file)

    result = read_docx(output_file)

    assert "John Smith" in result
    assert "Alice Brown" in result
    assert "+1 234 567 8910" not in result
    assert "+44 123 456 7890" not in result


def test_remove_phone_numbers_txt(tmp_path):
    input_file = tmp_path / "example.txt"
    output_file = tmp_path / "result.txt"

    input_file.write_text(
        "John Smith, +1 234 567 8910\nAlice Brown, +44 123 456 7890",
        encoding="utf-8")

    remove_phone_numbers_txt(input_file, output_file)

    result = output_file.read_text(encoding="utf-8")

    assert "John Smith" in result
    assert "Alice Brown" in result
    assert "+1 234 567 8910" not in result
    assert "+44 123 456 7890" not in result


def test_remove_phone_numbers_pdf(tmp_path):
    input_file = tmp_path / "example.pdf"
    output_file = tmp_path / "result.pdf"

    create_pdf(
        "John Smith, +1 234 567 8910\nAlice Brown, +44 123 456 7890",
        input_file)

    remove_phone_numbers_pdf(input_file, output_file)

    result = read_pdf(output_file)

    assert "John Smith" in result
    assert "Alice Brown" in result
    assert "234 567 8910" not in result
    assert "123 456 7890" not in result


def test_create_pdf(tmp_path):
    output_file = tmp_path / "result.pdf"

    create_pdf("Hello PDF", output_file)

    assert output_file.exists()

    document = pymupdf.open(output_file)
    text = document[0].get_text()
    document.close()

    assert "Hello PDF" in text
